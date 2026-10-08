"""MM-5B-PPT 的离线视频讲解到 PPTX 回归。

使用本机合成视频和固定讲解计划覆盖上游 HTML、关键帧复核、PPTX 回读、项目隔离与 API
下载。全程不调用模型和网络。
"""

from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from hashlib import sha256
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
from types import SimpleNamespace
from uuid import uuid4


BACKEND_ROOT = Path(__file__).resolve().parents[1]
VERIFY_ROOT = Path(tempfile.mkdtemp(prefix="agentflow_media_video_brief_presentation_"))
os.environ["AGENTFLOW_DATA_DIR"] = str(VERIFY_ROOT / "data")
os.environ["AGENTFLOW_OUTPUT_DIR"] = str(VERIFY_ROOT / "output")
os.environ["AGENTFLOW_MEDIA_VIDEO_BRIEF_OUTPUT_DIR"] = str(VERIFY_ROOT / "output" / "media_video_briefs")
os.environ["AGENTFLOW_MEDIA_VIDEO_PRESENTATION_OUTPUT_DIR"] = str(
    VERIFY_ROOT / "output" / "media_video_presentations"
)
os.environ["AGENTFLOW_DATABASE_PATH"] = str(VERIFY_ROOT / f"media-video-brief-presentation-{uuid4().hex}.db")
sys.path.insert(0, str(BACKEND_ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.api import media_agent as media_api  # noqa: E402
from app.api.tasks import _resolve_runtime_artifact_path  # noqa: E402
from app.database.task_repository import list_workflow_artifacts, list_workflow_tool_calls  # noqa: E402
from app.schemas.media_source import MediaTranscriptionSegmentInfo  # noqa: E402
from app.schemas.media_video_brief import (  # noqa: E402
    MediaVideoBriefModelChapter,
    MediaVideoBriefModelFact,
    MediaVideoBriefModelPlan,
    MediaVideoBriefPresentationRequest,
    MediaVideoBriefRequest,
)
from app.services.media_source_preparation import import_media_source_bytes, probe_media_source  # noqa: E402
from app.services.media_video_brief_delivery import (  # noqa: E402
    create_media_video_brief_queued_run,
    run_media_video_brief_task,
)
from app.services.media_video_brief_planning import MediaVideoBriefPlanningContext  # noqa: E402
from app.services.media_video_brief_presentation_delivery import (  # noqa: E402
    MediaVideoBriefPresentationError,
    create_media_video_brief_presentation_queued_run,
    resolve_media_video_brief_presentation_download_path,
    run_media_video_brief_presentation_task,
)
from app.services.media_workspace import create_media_project  # noqa: E402
from main import create_app  # noqa: E402


def _tool_paths() -> tuple[str, str]:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise RuntimeError("ffmpeg was not found on PATH")
    ffprobe = str(Path(ffmpeg).with_name("ffprobe.exe" if os.name == "nt" else "ffprobe"))
    if not Path(ffprobe).is_file():
        raise RuntimeError("ffprobe was not found next to ffmpeg")
    return ffmpeg, ffprobe


def _make_fixture(path: Path, ffmpeg: str) -> None:
    completed = subprocess.run(
        [
            ffmpeg,
            "-nostdin",
            "-hide_banner",
            "-loglevel",
            "error",
            "-y",
            "-f",
            "lavfi",
            "-i",
            "testsrc2=size=640x360:rate=25:duration=5",
            "-f",
            "lavfi",
            "-i",
            "sine=frequency=660:sample_rate=48000:duration=5",
            "-shortest",
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "aac",
            str(path),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    if completed.returncode != 0 or not path.is_file():
        raise RuntimeError("could not generate the synthetic video presentation fixture")


def _context(*, project_id: str, source_id: str, source_sha256: str, request: MediaVideoBriefRequest) -> MediaVideoBriefPlanningContext:
    return MediaVideoBriefPlanningContext(
        project_id=project_id,
        request=request,
        source_id=source_id,
        source_sha256=source_sha256,
        segments=(
            MediaTranscriptionSegmentInfo(sentence_id=10, text="系统先展示问题背景。", begin_ms=200, end_ms=1_100),
            MediaTranscriptionSegmentInfo(sentence_id=20, text="随后介绍核心处理流程。", begin_ms=1_150, end_ms=2_400),
            MediaTranscriptionSegmentInfo(sentence_id=30, text="最后说明交付结果与后续动作。", begin_ms=2_500, end_ms=4_300),
        ),
    )


def _model_plan() -> MediaVideoBriefModelPlan:
    return MediaVideoBriefModelPlan(
        action="brief",
        title="产品演示讲解",
        chapters=[
            MediaVideoBriefModelChapter(
                title="问题与处理流程",
                sentence_ids=[10, 20],
                facts=[MediaVideoBriefModelFact(text="视频先说明背景，再介绍处理流程。", sentence_ids=[10, 20])],
                layout="chapter",
                animation="appear",
            ),
            MediaVideoBriefModelChapter(
                title="交付结果",
                sentence_ids=[30],
                facts=[MediaVideoBriefModelFact(text="视频最后说明交付结果与后续动作。", sentence_ids=[30])],
                layout="summary",
                animation="fade",
            ),
        ],
    )


async def _fixed_planner(**_: object) -> MediaVideoBriefModelPlan:
    return _model_plan()


async def _run() -> dict[str, object]:
    ffmpeg, ffprobe = _tool_paths()
    fixture_path = VERIFY_ROOT / "fixture.mp4"
    _make_fixture(fixture_path, ffmpeg)
    project = create_media_project(title="Video presentation fixture")
    source_bytes = fixture_path.read_bytes()
    source = import_media_source_bytes(project_scope=project.project_id, filename="fixture.mp4", content=source_bytes)
    probe = probe_media_source(source_id=source.source_id, expected_project_scope=project.project_id, ffprobe_executable=ffprobe)
    assert probe.video_streams and probe.duration_seconds and 4.5 <= probe.duration_seconds <= 5.5

    brief_request = MediaVideoBriefRequest(
        transcription_task_id="task_media_transcription_0123456789ab",
        goal="整理为两章的可追溯视频讲解网页。",
    )
    context = _context(
        project_id=project.project_id,
        source_id=source.source_id,
        source_sha256=sha256(source_bytes).hexdigest(),
        request=brief_request,
    )
    brief_task_id = "task_media_video_brief_0123456789ab"
    create_media_video_brief_queued_run(task_id=brief_task_id, project_id=project.project_id, request=brief_request)
    brief_result = await run_media_video_brief_task(
        task_id=brief_task_id,
        project_id=project.project_id,
        request=brief_request,
        runtime=object(),
        planner=_fixed_planner,
        context_loader=lambda **_: context,
    )
    assert brief_result.status == "completed" and brief_result.plan is not None and brief_result.delivery is not None

    request = MediaVideoBriefPresentationRequest(video_brief_task_id=brief_task_id, confirmed=True)
    task_id = "task_media_video_presentation_0123456789ab"
    create_media_video_brief_presentation_queued_run(task_id=task_id, project_id=project.project_id, request=request)
    result = run_media_video_brief_presentation_task(task_id=task_id, project_id=project.project_id, request=request)
    assert result.status == "completed" and result.delivery is not None, result
    assert result.delivery.slide_count == 6 and result.delivery.embedded_keyframe_count == 2
    assert result.video_brief_task_id == brief_task_id and result.source_id == source.source_id
    path, filename = resolve_media_video_brief_presentation_download_path(project_id=project.project_id, task_id=task_id)
    assert filename == "video-brief.pptx" and path.is_file()
    opened = Presentation(path)
    assert len(opened.slides) == 6
    assert sum(1 for slide in opened.slides for shape in slide.shapes if hasattr(shape, "image")) >= 2
    artifacts = list_workflow_artifacts(task_id)
    calls = list_workflow_tool_calls(task_id)
    assert len(artifacts) == 1 and artifacts[0].mime_type.endswith("presentation")
    assert artifacts[0].metadata["video_brief_task_id"] == brief_task_id
    assert artifacts[0].metadata["model_used"] is False and artifacts[0].metadata["network_used"] is False
    assert _resolve_runtime_artifact_path(artifacts[0]) == path
    assert len(calls) == 1 and calls[0].request["model_used"] is False and calls[0].request["network_used"] is False

    other_project = create_media_project(title="Video presentation isolation fixture")
    try:
        resolve_media_video_brief_presentation_download_path(project_id=other_project.project_id, task_id=task_id)
    except MediaVideoBriefPresentationError:
        pass
    else:
        raise AssertionError("cross-project PPTX download must be rejected")

    with TestClient(create_app()) as client:
        missing_confirmation = client.post(
            f"/api/agents/media_agent/projects/{project.project_id}/video-brief-presentations/start",
            json={"video_brief_task_id": brief_task_id, "confirmed": False},
        )
        assert missing_confirmation.status_code == 422
        started = client.post(
            f"/api/agents/media_agent/projects/{project.project_id}/video-brief-presentations/start",
            json=request.model_dump(mode="json"),
        )
        assert started.status_code == 202, started.text
        api_task_id = started.json()["task_id"]
        api_result: dict[str, object] | None = None
        for _ in range(100):
            response = client.get(f"/api/agents/media_agent/video-brief-presentations/{api_task_id}/result")
            assert response.status_code == 200, response.text
            candidate = response.json()
            if candidate["status"] == "completed":
                api_result = candidate
                break
            assert candidate["status"] in {"pending", "running"}, candidate
            time.sleep(0.08)
        assert api_result is not None and api_result["delivery"]["embedded_keyframe_count"] == 2
        download = client.get(
            f"/api/agents/media_agent/projects/{project.project_id}/video-brief-presentations/{api_task_id}/download"
        )
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        assert len(download.content) > 4_096

    return {
        "ok": True,
        "network_used": False,
        "model_used": False,
        "slide_count": result.delivery.slide_count,
        "cases": ["upstream_binding", "keyframe_hash", "editable_pptx", "task_history_scope", "project_isolation", "api_start_poll_download"],
    }


def main() -> None:
    try:
        print(asyncio.run(_run()))
    finally:
        shutil.rmtree(VERIFY_ROOT, ignore_errors=True)


if __name__ == "__main__":
    main()
